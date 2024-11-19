import pymysql
import time
import pygetwindow as gw
from datetime import datetime
from pynput.mouse import Controller
from openpyxl import Workbook
import customtkinter as ctk
from tkinter import filedialog, messagebox
from tkcalendar import Calendar
import threading
from config import DATABASE_CONFIG

class DatabaseManager:
    def __init__(self):
        self.connection = self.get_connection()
        if self.connection:
            self.create_tables()

    def get_connection(self):
        try:
            connection = pymysql.connect(**DATABASE_CONFIG)
            if connection.open:
                print("Database connection established")
                return connection
            print("Database connection error")
            return None
        except pymysql.Error as e:
            print(f"Error connecting to MySQL: {e}")
            return None

    def create_tables(self):

        create_activity_log_table = """
        CREATE TABLE IF NOT EXISTS activity_log (
            id INT AUTO_INCREMENT PRIMARY KEY,
            user_id INT NOT NULL,
            date DATE NOT NULL,
            application VARCHAR(255) NOT NULL,
            time_spent INT NOT NULL,
            lag_points INT NOT NULL,
            pc_unlock_time INT NOT NULL,
            screen_saver_time INT NOT NULL,
            FOREIGN KEY (user_id) REFERENCES employees(user_id)
        )
        """
        try:
            with self.connection.cursor() as cursor:
                #cursor.execute(create_employees_table)
                cursor.execute(create_activity_log_table)
                self.connection.commit()
                print("Tables created or already exist")
        except pymysql.Error as e:
            print(f"Error creating tables: {e}")

    def fetch_employees(self):
        try:
            with self.connection.cursor() as cursor:
                cursor.execute("SELECT user_id, name FROM employees")
                return cursor.fetchall()
        except pymysql.Error as e:
            print(f"Error fetching employees: {e}")
            return []

    def fetch_activity_log(self, user_id, date):
        try:
            with self.connection.cursor() as cursor:
                cursor.execute("SELECT * FROM activity_log WHERE user_id = %s AND date = %s", (user_id, date))
                return cursor.fetchall()
        except pymysql.Error as e:
            print(f"Error fetching activity log: {e}")
            return []

def monitor_system(db, user_id):
    mouse = Controller()
    last_cursor_pos = mouse.position
    last_app = None
    last_unlock_time = time.time()
    lag_count = 0
    time_spent = 0
    screen_saver_time = 0

    while True:
        try:
            # Check active application
            current_app = gw.getActiveWindow().title if gw.getActiveWindow() else "Unknown"
            if current_app != last_app:
                if last_app is not None:
                    # Save the previous application's data
                    with db.connection.cursor() as cursor:
                        cursor.execute("""
                        INSERT INTO activity_log (user_id, date, application, time_spent, lag_points, pc_unlock_time, screen_saver_time)
                        VALUES (%s, %s, %s, %s, %s, %s, %s)
                        """, (user_id, datetime.now().date(), last_app, time_spent, lag_count, int(time.time() - last_unlock_time), screen_saver_time))
                        db.connection.commit()
                last_app = current_app
                time_spent = 0
                lag_count = 0

            # Check cursor stability (lag points)
            current_cursor_pos = mouse.position
            if current_cursor_pos == last_cursor_pos:
                lag_count += 1
            last_cursor_pos = current_cursor_pos

            # Simulate unlock time tracking (add custom logic if required)
            last_unlock_time = time.time()  # Update last unlock time

            time.sleep(1)
            time_spent += 1
        except Exception as e:
            print(f"Error: {e}")

def generate_report(db, user_ids, date):
    try:
        wb = Workbook()
        for user_id in user_ids:
            with db.connection.cursor() as cursor:
                cursor.execute("SELECT name FROM employees WHERE user_id = %s", (user_id,))
                employee_name = cursor.fetchone()[0]
                cursor.execute("SELECT * FROM activity_log WHERE user_id = %s AND date = %s", (user_id, date))
                data = cursor.fetchall()

            ws = wb.create_sheet(title=employee_name)
            ws.append(["Date", "Application", "Time Spent (s)", "Lag Points", "PC Unlock Time (s)", "Screen Saver Time (s)"])
            for row in data:
                ws.append(row[2:])  # Skip user_id and id columns

        # Remove the default sheet created by openpyxl
        wb.remove(wb["Sheet"])

        # Ask user where to save the file
        report_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
            title="Save Report As"
        )
        if report_path:
            wb.save(report_path)
            print(f"Report saved at {report_path}")
        else:
            print("Save operation canceled.")
    except pymysql.Error as e:
        print(f"Error generating report: {e}")

def create_gui(db):
    def on_generate_report():
        selected_date = calendar.get_date().replace("-", "")  # Format date as yyyymmdd
        selected_user_ids = [user_id for user_id, var in employee_vars.items() if var.get()]
        if not selected_user_ids:
            messagebox.showwarning("No Selection", "Please select at least one employee.")
            return
        generate_report(db, selected_user_ids, selected_date)

    app = ctk.CTk()
    app.title("Employee Activity Tracker")
    app.geometry("500x400")

    date_label = ctk.CTkLabel(app, text="Select Date:")
    date_label.pack(pady=10)

    # Calendar widget for date selection
    calendar = Calendar(app, selectmode='day', date_pattern='yyyy-mm-dd')
    calendar.pack(pady=10)

    # Employee selection
    employee_label = ctk.CTkLabel(app, text="Select Employees to Monitor:")
    employee_label.pack(pady=10)

    employee_vars = {}
    employees = db.fetch_employees()
    for user_id, name in employees:
        var = ctk.BooleanVar()
        chk = ctk.CTkCheckBox(app, text=name, variable=var)
        chk.pack()
        employee_vars[user_id] = var

    generate_button = ctk.CTkButton(app, text="Generate Report", command=on_generate_report)
    generate_button.pack(pady=20)

    app.mainloop()

if __name__ == "__main__":
    db = DatabaseManager()
    create_gui(db)



# import sqlite3
# import pymysql
# import time
# import pygetwindow as gw
# from datetime import datetime
# from pynput.mouse import Controller
# from openpyxl import Workbook
# import customtkinter as ctk
# from tkinter import filedialog
# from tkcalendar import Calendar
# import threading
# from config import DATABASE_CONFIG


# class DatabaseManager:
#     def __init__(self):
#         pass

#     def get_connection(self):
#         try:
#             connection = pymysql.connect(**DATABASE_CONFIG)
#             if connection.open:
#                 print("Database connection established")
#                 return connection
#             print("Database connection error")
#             return None
#         except pymysql.Error as e:
#             print(f"Error connecting to MySQL: {e}")
#             return None
        


# # Monitor system activity
# def monitor_system():
#     try :
#         conn = self.db.get
#         cursor = conn.cursor()
#     except Error as e:
#         print("Error",str(e))

#     mouse = Controller()
#     last_cursor_pos = mouse.position
#     last_app = None
#     last_unlock_time = time.time()
#     lag_count = 0
#     time_spent = 0
#     screen_saver_time = 0

#     while True:
#         try:
#             # Check active application
#             current_app = gw.getActiveWindow().title if gw.getActiveWindow() else "Unknown"
#             if current_app != last_app:
#                 if last_app is not None:
#                     # Save the previous application's data
#                     cursor.execute("""
#                     INSERT INTO activity_log (date, application, time_spent, lag_points, pc_unlock_time, screen_saver_time)
#                     VALUES (?, ?, ?, ?, ?, ?)
#                     """, (datetime.now().date(), last_app, time_spent, lag_count, int(time.time() - last_unlock_time), screen_saver_time))
#                     conn.commit()
#                 last_app = current_app
#                 time_spent = 0
#                 lag_count = 0

#             # Check cursor stability (lag points)
#             current_cursor_pos = mouse.position
#             if current_cursor_pos == last_cursor_pos:
#                 lag_count += 1
#             last_cursor_pos = current_cursor_pos

#             # Simulate unlock time tracking (add custom logic if required)
#             last_unlock_time = time.time()  # Update last unlock time

#             time.sleep(1)
#             time_spent += 1
#         except Exception as e:
#             print(f"Error: {e}")

# # Generate Excel report
# def generate_report(date):
#     conn = sqlite3.connect(DB_FILE)
#     cursor = conn.cursor()
#     cursor.execute("SELECT * FROM activity_log WHERE date = ?", (date,))
#     data = cursor.fetchall()
#     conn.close()

#     wb = Workbook()
#     ws = wb.active
#     ws.append(["Date", "Application", "Time Spent (s)", "Lag Points", "PC Unlock Time (s)", "Screen Saver Time (s)"])
#     for row in data:
#         ws.append(row[1:])
    
#     # Ask user where to save the file
#     report_path = filedialog.asksaveasfilename(
#         defaultextension=".xlsx",
#         filetypes=[("Excel Files", "*.xlsx")],
#         title="Save Report As"
#     )
#     if report_path:
#         wb.save(report_path)
#         print(f"Report saved at {report_path}")
#     else:
#         print("Save operation canceled.")

# # GUI with CustomTkinter and Calendar
# def create_gui():
#     def on_generate_report():
#         selected_date = calendar.get_date()
#         generate_report(selected_date)
    
#     app = ctk.CTk()
#     app.title("Activity Tracker Report")
#     app.geometry("400x400")

#     date_label = ctk.CTkLabel(app, text="Select Date:")
#     date_label.pack(pady=10)

#     # Calendar widget for date selection
#     calendar = Calendar(app, selectmode='day', date_pattern='yyyy-mm-dd')
#     calendar.pack(pady=10)

#     generate_button = ctk.CTkButton(app, text="Generate Report", command=on_generate_report)
#     generate_button.pack(pady=20)

#     app.mainloop()

# if __name__ == "_main_":
#     setup_database()
#     threading.Thread(target=monitor_system, daemon=True).start()
#     create_gui()